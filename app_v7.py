"""
AI AUTONOMOUS COGNITIVE ENGINE — Deep Research & Long Horizon Tasks
Streamlit Web Interface — v7
New in v7:
  - RAG: PDF/URL/text document upload + retrieval
  - Continuous conversation with chat history
  - Critic evaluation shown after final report
  - Download as PDF and Word (.docx)
"""
import os
import re
import sys
import time
import io
import streamlit as st

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "src"))

# ------------------------------------------------------------------ #
#  Page config                                                        #
# ------------------------------------------------------------------ #
st.set_page_config(
    page_title="AI Autonomous Cognitive Deep-Research Engine",
    page_icon="\U0001f9e0",
    layout="wide",
    initial_sidebar_state="expanded",
)

# ------------------------------------------------------------------ #
#  Load secrets / env                                                 #
# ------------------------------------------------------------------ #
def _load_secrets():
    for key in [
        "ANTHROPIC_API_KEY", "ANTHROPIC_MODEL",
        "GEMINI_API_KEY",    "GEMINI_MODEL", "GEMINI_ROLES",
        "OPENROUTER_API_KEY", "OPENROUTER_MODEL",
        "TAVILY_API_KEY",
    ]:
        if key in st.secrets and not os.environ.get(key):
            os.environ[key] = st.secrets[key]

try:
    _load_secrets()
except Exception:
    pass

try:
    from dotenv import load_dotenv
    load_dotenv()
except ImportError:
    pass

# ------------------------------------------------------------------ #
#  Helper: strip tool-call artifacts                                  #
# ------------------------------------------------------------------ #
def _strip_artifacts(text: str) -> str:
    if not text:
        return text
    text = re.sub(r'<\|tool_call_start\|>.*?<\|tool_call_end\|>', '', text, flags=re.DOTALL)
    text = re.sub(r'\[(?:google|search|fetch)\(.*?\)(?:,\s*(?:google|search|fetch)\(.*?\))*\]', '', text, flags=re.DOTALL)
    return text.strip()

# ------------------------------------------------------------------ #
#  Helper: generate Word (.docx) bytes from markdown text             #
# ------------------------------------------------------------------ #
def _to_docx(text: str) -> bytes:
    try:
        from docx import Document as DocxDocument
        from docx.shared import Pt
        doc = DocxDocument()
        doc.add_heading("AI Research Report", 0)
        for line in text.split("\n"):
            stripped = line.strip()
            if stripped.startswith("### "):
                doc.add_heading(stripped[4:], level=3)
            elif stripped.startswith("## "):
                doc.add_heading(stripped[3:], level=2)
            elif stripped.startswith("# "):
                doc.add_heading(stripped[2:], level=1)
            elif stripped.startswith("- ") or stripped.startswith("* "):
                doc.add_paragraph(stripped[2:], style="List Bullet")
            elif stripped:
                doc.add_paragraph(stripped)
        buf = io.BytesIO()
        doc.save(buf)
        return buf.getvalue()
    except Exception:
        return text.encode("utf-8")

# ------------------------------------------------------------------ #
#  Helper: generate PDF bytes from markdown text                      #
# ------------------------------------------------------------------ #
def _to_pdf(text: str) -> bytes:
    try:
        from fpdf import FPDF
        pdf = FPDF()
        pdf.set_auto_page_break(auto=True, margin=15)
        pdf.add_page()
        pdf.set_font("Helvetica", size=11)
        for line in text.split("\n"):
            stripped = line.strip()
            if stripped.startswith("# "):
                pdf.set_font("Helvetica", "B", 16)
                pdf.multi_cell(0, 8, stripped[2:])
                pdf.set_font("Helvetica", size=11)
            elif stripped.startswith("## "):
                pdf.set_font("Helvetica", "B", 13)
                pdf.multi_cell(0, 7, stripped[3:])
                pdf.set_font("Helvetica", size=11)
            elif stripped.startswith("### "):
                pdf.set_font("Helvetica", "B", 11)
                pdf.multi_cell(0, 6, stripped[4:])
                pdf.set_font("Helvetica", size=11)
            elif stripped.startswith("- ") or stripped.startswith("* "):
                pdf.multi_cell(0, 6, "  - " + stripped[2:])
            elif stripped:
                pdf.multi_cell(0, 6, stripped)
            else:
                pdf.ln(3)
        return pdf.output(dest="S").encode("latin-1")
    except Exception:
        return text.encode("utf-8")

# ------------------------------------------------------------------ #
#  Session state init                                                 #
# ------------------------------------------------------------------ #
if "chat_history" not in st.session_state:
    st.session_state.chat_history = []
if "conversation_context" not in st.session_state:
    st.session_state.conversation_context = ""
if "last_critic" not in st.session_state:
    st.session_state.last_critic = None

# ------------------------------------------------------------------ #
#  Sidebar                                                            #
# ------------------------------------------------------------------ #
with st.sidebar:
    st.image("https://cdn-icons-png.flaticon.com/512/4616/4616013.png", width=64)
    st.title("\U0001f9e0 AI Autonomous Cognitive Deep-Research Engine")
    st.caption("Deep-Research \u00b7 Long Horizon Tasks \u00b7 Autonomous Multi-Agent")
    st.divider()

    st.markdown("### \U0001f9e0 Agent Architecture")
    st.markdown(
        "| Agent | Role |\n"
        "|---|---|\n"
        "| \U0001f4cc Planner | Decomposes query into tasks |\n"
        "| \U0001f50d Search | Live web search & evidence |\n"
        "| \U0001f52c Research | Deep analysis & synthesis |\n"
        "| \U0001f4ca Math | Calculations & proofs |\n"
        "| \U0001f4bb Coding | Code generation & debug |\n"
        "| \u270d\ufe0f Writer | Final report synthesis |\n"
        "| \u2696\ufe0f Critic | Quality gate + retry loop |"
    )
    st.divider()

    # ---- RAG: Knowledge Base ---- #
    st.markdown("### \U0001f4da Knowledge Base (RAG)")
    st.caption("Upload documents — agents will reference them during research.")

    uploaded_file = st.file_uploader(
        "Upload PDF or TXT",
        type=["pdf", "txt"],
        help="Agent will use this as a primary source"
    )
    if uploaded_file is not None:
        try:
            from ai_deep_agent.memory.vector_store import ingest_pdf, ingest_text
            import tempfile
            suffix = "." + uploaded_file.name.split(".")[-1]
            with tempfile.NamedTemporaryFile(delete=False, suffix=suffix) as f:
                f.write(uploaded_file.read())
                tmp_path = f.name
            with st.spinner("Indexing document..."):
                if suffix == ".pdf":
                    count = ingest_pdf(tmp_path)
                else:
                    content = open(tmp_path, encoding="utf-8", errors="ignore").read()
                    count = ingest_text(content, uploaded_file.name)
            st.success(f"\u2705 {count} chunks indexed from {uploaded_file.name}")
        except Exception as e:
            st.error(f"Upload failed: {e}")

    url_input = st.text_input("\U0001f310 Or paste a URL to index")
    if st.button("Index URL") and url_input.strip():
        try:
            from ai_deep_agent.memory.vector_store import ingest_url
            with st.spinner("Fetching and indexing..."):
                count = ingest_url(url_input.strip())
            st.success(f"\u2705 {count} chunks indexed!")
        except Exception as e:
            st.error(f"URL indexing failed: {e}")

    col_a, col_b = st.columns(2)
    with col_a:
        if st.button("\U0001f5d1\ufe0f Clear KB", use_container_width=True):
            try:
                from ai_deep_agent.memory.vector_store import clear_vector_store
                clear_vector_store()
                st.success("Cleared!")
            except Exception as e:
                st.error(str(e))
    with col_b:
        if st.button("\U0001f504 New Chat", use_container_width=True):
            st.session_state.chat_history = []
            st.session_state.conversation_context = ""
            st.session_state.last_critic = None
            st.rerun()

    st.divider()
    model = (
        os.environ.get("OPENROUTER_MODEL")
        or os.environ.get("GEMINI_MODEL")
        or os.environ.get("ANTHROPIC_MODEL")
        or "Kya Farak Padta Hai, Chal toh raha hai na?"
    )
    st.markdown("### \u2699\ufe0f Config")
    st.code(f"Model: {model}")
    st.divider()
    st.caption("\U0001f468\u200d\U0001f4bb Built by Manishka | Placement Project")

# ------------------------------------------------------------------ #
#  Main UI Header                                                     #
# ------------------------------------------------------------------ #
st.markdown(
    "<h1 style='text-align:center;'>\U0001f9e0 AI Autonomous Cognitive Deep-Research Engine</h1>"
    "<p style='text-align:center; color:gray;'>"
    "Ask anything \u2014 the agent plans, researches, reasons, and writes a full report."
    "</p>",
    unsafe_allow_html=True,
)
st.divider()

# ------------------------------------------------------------------ #
#  Chat history display                                               #
# ------------------------------------------------------------------ #
for msg in st.session_state.chat_history:
    with st.chat_message(msg["role"]):
        if msg["role"] == "user":
            st.markdown(msg["content"])
        else:
            st.markdown(msg["content"])
            if msg.get("critic"):
                critic = msg["critic"]
                decision = critic.get("decision", "pass")
                if decision == "pass":
                    st.success("\u2696\ufe0f **Critic Verdict: PASS** \u2014 " + critic.get("reason", "")[:150])
                else:
                    st.error("\u2696\ufe0f **Critic Verdict: FAIL** \u2014 " + critic.get("reason", "")[:150])

# ------------------------------------------------------------------ #
#  Chat input                                                         #
# ------------------------------------------------------------------ #
query = st.chat_input("\U0001f50d Ask anything...")

if query and query.strip():

    missing = [k for k in ["TAVILY_API_KEY"]
               if not os.environ.get(k, "").strip()]
    if missing:
        st.error(f"\U0001f534 Missing API keys: {', '.join(missing)}. Add to Streamlit Secrets or .env")
        st.stop()

    # Show user message
    with st.chat_message("user"):
        st.markdown(query)
    st.session_state.chat_history.append({"role": "user", "content": query})

    # Build conversation-aware query
    full_query = query.strip()
    if st.session_state.conversation_context:
        full_query = (
            f"Previous research context:\n{st.session_state.conversation_context[:800]}\n\n"
            f"New question (build on above context):\n{query.strip()}"
        )

    from ai_deep_agent.graph.builder import app
    from ai_deep_agent.state.state   import AgentState
    import ai_deep_agent.display.console as ui

    with st.chat_message("assistant"):

        # ---- Live Process Display ---- #
        st.markdown("### \u26a1 Live Agent Process")
        process_box = st.empty()
        live_log = []

        ICONS = {
            "search":   "\U0001f50d",
            "research": "\U0001f52c",
            "math":     "\U0001f4ca",
            "coding":   "\U0001f4bb",
            "writer":   "\u270d\ufe0f",
        }

        def _refresh():
            with process_box.container():
                for line in live_log:
                    st.markdown(line, unsafe_allow_html=True)

        def _log(msg: str):
            live_log.append(msg)
            _refresh()

        # ---- Monkey-patch ui.* ---- #
        _orig_planner_done = ui.planner_done
        def _st_planner_done(todos):
            _orig_planner_done(todos)
            _log("")
            _log(f"**\U0001f4cc Planning complete \u2014 {len(todos)} task(s) created:**")
            for t in todos:
                _log(f"&nbsp;&nbsp;&nbsp;&nbsp;`{t['id']}.` {t['task']}")
            _log("---")
        ui.planner_done = _st_planner_done

        _orig_task_start = ui.task_start
        def _st_task_start(task_id, total, task, worker):
            _orig_task_start(task_id, total, task, worker)
            icon = ICONS.get(worker, "\u2699\ufe0f")
            _log(f"{icon} **Task {task_id}/{total}** \u2192 `{worker.capitalize()} Agent` \u2014 *{task[:80]}...*")
        ui.task_start = _st_task_start

        _orig_task_accepted = ui.task_accepted
        def _st_task_accepted(task_id, worker, retries):
            _orig_task_accepted(task_id, worker, retries)
            _log(f"&nbsp;&nbsp;&nbsp;&nbsp;\u2705 Task {task_id} complete")
        ui.task_accepted = _st_task_accepted

        _orig_worker_result = ui.worker_result
        def _st_worker_result(worker, result, attempt):
            _orig_worker_result(worker, result, attempt)
            label = f" (attempt #{attempt + 1})" if attempt > 0 else ""
            _log(f"\u270d\ufe0f **Writer Agent{label}** \u2014 synthesising final report...")
        ui.worker_result = _st_worker_result

        # Store last critic result for display after report
        _last_critic_store = [None]

        _orig_critic_pass = ui.critic_pass
        def _st_critic_pass(reason, strengths):
            _orig_critic_pass(reason, strengths)
            _log(f"\u2705 **Critic \u2014 PASS:** {reason[:120]}")
            for s in (strengths or [])[:3]:
                _log(f"&nbsp;&nbsp;&nbsp;&nbsp;\u2728 {s}")
            _last_critic_store[0] = {"decision": "pass", "reason": reason, "strengths": strengths or []}
        ui.critic_pass = _st_critic_pass

        _orig_critic_fail = ui.critic_fail
        def _st_critic_fail(reason, missing, instructions):
            _orig_critic_fail(reason, missing, instructions)
            _log(f"\u274c **Critic \u2014 FAIL:** {reason[:120]}")
            for m in (missing or [])[:4]:
                _log(f"&nbsp;&nbsp;&nbsp;&nbsp;\u2022 Missing: {m}")
            for fix in (instructions or [])[:2]:
                _log(f"&nbsp;&nbsp;&nbsp;&nbsp;\U0001f527 Fix: {fix[:100]}")
            _last_critic_store[0] = {"decision": "fail", "reason": reason}
        ui.critic_fail = _st_critic_fail

        _orig_retry = ui.retry_notice
        def _st_retry(attempt, max_retries, feedback):
            _orig_retry(attempt, max_retries, feedback)
            _log(f"\U0001f504 **Retry {attempt}/{max_retries}** \u2014 sending feedback to Writer...")
            _log(f"&nbsp;&nbsp;&nbsp;&nbsp;*{feedback[:150]}*")
        ui.retry_notice = _st_retry

        _orig_max_retries = ui.max_retries_hit
        def _st_max_retries(task_id):
            _orig_max_retries(task_id)
            _log("\u26a0\ufe0f **Max retries reached** \u2014 using best available output")
        ui.max_retries_hit = _st_max_retries

        # ---- Kick off ---- #
        _log("\U0001f9e0 **Agent started** \u2014 analysing your query...")
        _log(f"> *{query.strip()[:120]}*")
        _log("---")
        _log("\u23f3 **Planning** \u2014 breaking query into research tasks...")

        initial: AgentState = {
            "messages":           [],
            "user_query":         full_query,
            "todos":              [],
            "planning_complete":  False,
            "completed_tasks":    [],
            "final_answer":       "",
            "writer_used":        False,
            "retry_counts":       {},
            "execution_log":      [],
            "sources":            [],
            "retrieved_context":  "",
            "chat_history":       st.session_state.chat_history,
        }

        start   = time.time()
        result  = app.invoke(initial)
        elapsed = time.time() - start

        _log("---")
        _log(f"\U0001f3c1 **Done in {elapsed:.1f}s** \u2014 report ready below \u2193")

        # ---- Final Report ---- #
        st.divider()
        st.success(f"\u2705 Completed in {elapsed:.1f}s")
        st.markdown("## \U0001f4dd Final Report")
        st.divider()

        final_report = _strip_artifacts(result.get("final_answer", ""))
        completed    = result.get("completed_tasks", [])

        if final_report:
            st.markdown(final_report)
        elif completed:
            st.markdown(_strip_artifacts(completed[-1].get("result", "*No output generated.*")))
        else:
            st.warning("\u26a0\ufe0f No output was generated. Try a different query or check your API keys.")

        # ---- Critic Verdict after report ---- #
        critic_result = _last_critic_store[0]
        if critic_result:
            st.divider()
            st.markdown("### \u2696\ufe0f Critic Evaluation")
            if critic_result["decision"] == "pass":
                st.success(f"\u2705 **PASS** \u2014 {critic_result.get('reason', '')[:200]}")
                for s in critic_result.get("strengths", [])[:5]:
                    st.markdown(f"&nbsp;&nbsp;\u2728 {s}")
            else:
                st.error(f"\u274c **FAIL** \u2014 {critic_result.get('reason', '')[:200]}")
            st.session_state.last_critic = critic_result

        # ---- Sources ---- #
        sources = result.get("sources", [])
        if sources:
            st.divider()
            with st.expander("\U0001f517 Sources", expanded=False):
                seen = set()
                for s in sources:
                    url = s.get("url", "")
                    if url and url not in seen:
                        seen.add(url)
                        title = s.get("title", url)
                        st.markdown(f"- [{title}]({url})")

        # ---- Execution Summary ---- #
        if completed:
            with st.expander("\U0001f4ca Execution Summary", expanded=False):
                worker_icons = {
                    "search": "\U0001f50d", "research": "\U0001f52c",
                    "math":   "\U0001f4ca", "coding":   "\U0001f4bb", "writer": "\u270d\ufe0f",
                }
                for task in completed:
                    worker  = task.get("worker", "")
                    icon    = worker_icons.get(worker, "\u2699\ufe0f")
                    tid     = task.get("id", "?")
                    tdesc   = task.get("task", "")[:80]
                    verdict = (task.get("review") or {}).get("decision", "pass")
                    badge   = "\u2705" if verdict == "pass" else "\U0001f504"
                    retries = task.get("retries", 0)
                    r_tag   = f" ({retries} retr{'y' if retries==1 else 'ies'})" if retries else ""
                    st.markdown(f"{badge} **Task {tid}** {icon} `{worker.capitalize()}`{r_tag} \u2014 {tdesc}...")

        # ---- Download: PDF + Word ---- #
        report_text = final_report or (completed[-1].get("result", "") if completed else "")
        if report_text:
            st.divider()
            st.markdown("#### \u2b07\ufe0f Download Report")
            dl1, dl2 = st.columns(2)
            with dl1:
                st.download_button(
                    label="\U0001f4c4 Download as Word (.docx)",
                    data=_to_docx(report_text),
                    file_name="ai_research_report.docx",
                    mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    use_container_width=True,
                )
            with dl2:
                st.download_button(
                    label="\U0001f4cb Download as PDF",
                    data=_to_pdf(report_text),
                    file_name="ai_research_report.pdf",
                    mime="application/pdf",
                    use_container_width=True,
                )

        # ---- Update conversation context ---- #
        st.session_state.conversation_context = (
            f"Q: {query.strip()}\nSummary: {report_text[:600]}"
        )
        st.session_state.chat_history.append({
            "role":    "assistant",
            "content": final_report or report_text,
            "critic":  critic_result,
        })
